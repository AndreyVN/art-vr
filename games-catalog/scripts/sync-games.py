# -*- coding: utf-8 -*-
"""Синхронизирует games.json с вкладкой «Игры» таблицы клуба ART-VR.

Как читается таблица:
  * основной список — игры на ПК: красная заливка названия — не работает, остальные — работают;
    пометка «- пока нету» в названии — ещё не установлена;
  * «Игры на ПС:» (приставка PlayStation 5) и «Настольные игры:» — пока в каталог
    не входят, скрипт их пропускает;
  * «Игры на шлемах:» — работают на автономных шлемах, все однопользовательские;
  * «Тестировали, но не понравились:» — статус rejected (или broken, если «Не работает»).

Что делает скрипт:
  * обновляет статус, платформы, игроков, возраст, время, жанр, зоны, рейтинг, ролик
    и примечание из таблицы;
  * у игр с "locked": true статус не трогает — это решения клуба (см. поле note);
  * поля из "locked_fields" не перезаписывает, а предупреждает, если в таблице другое;
  * игру, которой больше нет в таблице, помечает статусом not_in_club;
  * строку таблицы, которой нет в списке, добавляет черновиком (status: draft) —
    допишите описание, и при следующем запуске рабочая игра опубликуется;
  * ничего не удаляет; в конце пересобирает games-data.ts (scripts/generate-games.py).

Запуск:
  python scripts/sync-games.py                      # таблица по умолчанию
  python scripts/sync-games.py --xlsx "путь.xlsx"   # другая выгрузка
  python scripts/sync-games.py --dry-run            # только показать изменения
"""
import argparse
import datetime
import io
import json
import os
import re
import subprocess
import sys
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GAMES_JSON = os.path.join(ROOT, 'games.json')
DEFAULT_XLSX = os.path.join(os.path.dirname(os.path.dirname(ROOT)), 'ART-VR (2).xlsx')
SHEET = 'Игры'

RED = 'FFFF0000'

# Колонки вкладки «Игры».
(COL_NAME, COL_PLAYERS, COL_GENRE, COL_AGE, COL_TIME, COL_RATING, COL_ZONES,
 COL_REVIEW, COL_NOTE) = 1, 2, 3, 4, 5, 6, 7, 11, 12

BLOCKS = [
    ('игры на пс', 'ps5'),
    ('игры на шлемах', 'headset'),
    ('настольные игры', 'board'),
    ('тестировали, но не понравились', 'rejected'),
]

# Жанр из таблицы -> раздел каталога (для новых игр; у существующих раздел не меняется).
RAW_GENRE = {
    'pvp': 'shooter', 'оборона': 'shooter', 'оборона/pvp': 'shooter', 'оборона, pvp': 'shooter',
    'pvp/прохождение': 'shooter', 'шутер': 'shooter', 'стрелялка': 'shooter', 'бои': 'shooter',
    'дуэль': 'shooter', 'музыкальная': 'rhythm', 'музыкальная шутер': 'rhythm', 'музыка': 'rhythm',
    'спорт': 'sport', 'паркур': 'sport', 'квест': 'quest', 'логическая': 'quest',
    'прохождение': 'quest', 'симулятор': 'sim', 'готовка': 'sim', 'гонки': 'sim',
    'песочница': 'sim', 'выживание': 'sim', 'настолка': 'party', 'мафия': 'party',
    'хоррор': 'horror', 'смотреть': 'relax', 'плавать': 'relax', 'рисовалка': 'relax',
}

YOUTUBE_ID = re.compile(r'(?:v=|youtu\.be/|embed/|shorts/)([A-Za-z0-9_-]{11})')


def cell(value):
    # В таблице диапазоны вида «1-5» местами сохранились как даты (1 мая).
    if isinstance(value, datetime.datetime):
        return '%d-%d' % (value.day, value.month)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return '' if value is None else str(value).strip()


def norm(name):
    s = (name or '').lower().replace('’', "'").replace('&', ' and ')
    s = s.replace('а', 'a')  # «FNАF» в таблице набран с русской «А»
    s = re.sub(r'\((демо|beta|расширение)\)', '', s)
    s = re.sub(r'[^a-z0-9а-я]+', ' ', s)
    s = re.sub(r'\b(vr|the|a)\b', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()


def slugify(name):
    s = unicodedata.normalize('NFKD', name)
    s = s.replace("'", '').replace('’', '').replace('+', '-plus')
    s = re.sub(r'[^a-zA-Z0-9]+', '-', s).strip('-').lower()
    return s or 'game'


def clean_players(value):
    value = value.replace('?', '').strip()
    if not re.search(r'\d', value):
        return ''
    return re.sub(r'\s*,\s*', ', ', value)


def clean_age(value):
    m = re.match(r'\d+', value)
    return m.group(0) if m else ''


def youtube_url(value):
    m = YOUTUBE_ID.search(value or '')
    return 'https://www.youtube.com/watch?v=' + m.group(1) if m else ''


def read_sheet(path):
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    block = 'pc'
    rows = []
    for r in range(2, ws.max_row + 1):
        name = cell(ws.cell(r, COL_NAME).value)
        if not name:
            continue
        low = name.lower()
        header = next((b for prefix, b in BLOCKS if low.startswith(prefix)), None)
        if header:
            block = header
            continue
        if block in ('ps5', 'board'):
            continue
        fill = ws.cell(r, COL_NAME).fill
        rgb = fill.fgColor.rgb if fill is not None and fill.fill_type not in (None, 'none') else ''
        note = cell(ws.cell(r, COL_NOTE).value)
        missing = re.search(r'\s+-\s*пока нету\s*$', name, re.I)
        if missing:
            name = name[:missing.start()].strip()
            note = '; '.join(x for x in (note, 'пока нету') if x)
        if block == 'pc':
            status = 'not_installed' if missing else ('broken' if rgb == RED else 'works')
        elif block == 'headset':
            status = 'works'
        else:
            status = 'broken' if 'не работает' in note.lower() else 'rejected'
        players = cell(ws.cell(r, COL_PLAYERS).value)
        notes = [note]
        if players and not clean_players(players):
            notes.append(players)
        rows.append({
            'row': r, 'name': name, 'block': block, 'status': status,
            # На шлемах все игры однопользовательские (пометка в заголовке блока).
            'players': clean_players(players) or ('1' if block == 'headset' else ''),
            'genre': cell(ws.cell(r, COL_GENRE).value),
            'age': clean_age(cell(ws.cell(r, COL_AGE).value)),
            'playtime': cell(ws.cell(r, COL_TIME).value),
            'rating': cell(ws.cell(r, COL_RATING).value),
            'zones': cell(ws.cell(r, COL_ZONES).value),
            'youtube': youtube_url(cell(ws.cell(r, COL_REVIEW).value)),
            'note': '; '.join(x for x in notes if x),
        })
    return rows


def sheet_verdict(rows):
    """Статус и платформы игры по всем её строкам таблицы."""
    blocks = {x['block'] for x in rows}
    guess = (['pc'] if 'pc' in blocks else []) + (['standalone'] if 'headset' in blocks else [])
    rejected = [x for x in rows if x['block'] == 'rejected']
    if rejected:
        return rejected[0]['status'], guess
    pc = [x for x in rows if x['block'] == 'pc']
    pc_status = pc[0]['status'] if pc else None
    platforms = []
    if pc_status == 'works':
        platforms.append('pc')
    if 'headset' in blocks:
        platforms.append('standalone')
    if platforms:
        return 'published', platforms
    return pc_status or 'broken', guess


def publishable(game):
    return bool(game.get('description'))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--xlsx', default=DEFAULT_XLSX)
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    data = json.load(io.open(GAMES_JSON, encoding='utf-8'))
    games = data['games']
    rows = read_sheet(args.xlsx)

    index = {}
    for g in games:
        for name in [g['title']] + g.get('sheet_names', []):
            index.setdefault(norm(name), g)

    matched = {}
    fresh = {}
    for r in rows:
        g = index.get(norm(r['name']))
        if g is not None:
            matched.setdefault(g['slug'], []).append(r)
        else:
            fresh.setdefault(norm(r['name']), []).append(r)

    log = []
    for g in games:
        g_rows = matched.get(g['slug'])
        if not g_rows:
            if g['status'] != 'not_in_club':
                if g.get('locked'):
                    log.append('! %s: нет в таблице, статус %s закреплён' % (g['title'], g['status']))
                else:
                    log.append('- %s: нет в таблице, %s -> not_in_club' % (g['title'], g['status']))
                    g['status'] = 'not_in_club'
            continue

        status, platforms = sheet_verdict(g_rows)
        if status == 'published' and not publishable(g):
            status = 'draft'
        if platforms and platforms != g.get('platforms'):
            log.append('~ %s: платформы %s -> %s' % (g['title'], g.get('platforms'), platforms))
            g['platforms'] = platforms
        if not g.get('locked') and status != g['status']:
            log.append('~ %s: статус %s -> %s' % (g['title'], g['status'], status))
            g['status'] = status

        # Сначала строки основного списка: там полные данные, в блоке шлемов — только зоны.
        ordered = sorted(g_rows, key=lambda x: x['block'] != 'pc')
        for key in ('players', 'age', 'playtime', 'youtube'):
            value = next((x[key] for x in ordered if x[key]), '')
            if not value or value == g.get(key):
                continue
            if key in g.get('locked_fields', []):
                log.append('! %s: %s в таблице %r, в списке закреплено %r — поправьте таблицу' % (
                    g['title'], key, value, g.get(key, '')))
                continue
            log.append('~ %s: %s %r -> %r' % (g['title'], key, g.get(key, ''), value))
            g[key] = value
        g['sheet_genre'] = next((x['genre'] for x in ordered if x['genre']), '')
        g['sheet_rating'] = next((x['rating'] for x in ordered if x['rating']), '')
        g['sheet_zones'] = '; '.join(
            '%s: %s' % ({'pc': 'ПК', 'headset': 'шлемы'}.get(x['block'], x['block']), x['zones'])
            for x in ordered if x['zones'] and x['block'] in ('pc', 'headset'))
        g['sheet_note'] = '; '.join(dict.fromkeys(x['note'] for x in ordered if x['note']))
        names = g.setdefault('sheet_names', [])
        for x in g_rows:
            if x['name'] not in names:
                names.append(x['name'])

    slugs = {g['slug'] for g in games}
    for key, f_rows in fresh.items():
        name = f_rows[0]['name']
        status, platforms = sheet_verdict(f_rows)
        slug = slugify(name)
        while slug in slugs:
            slug += '-2'
        slugs.add(slug)
        genre_raw = next((x['genre'] for x in f_rows if x['genre']), '')
        games.append({
            'title': name,
            'slug': slug,
            'status': 'draft' if status == 'published' else status,
            'locked': False,
            'platforms': platforms,
            'genre': RAW_GENRE.get(genre_raw.lower(), 'quest'),
            'players': next((x['players'] for x in f_rows if x['players']), ''),
            'age': next((x['age'] for x in f_rows if x['age']), ''),
            'playtime': next((x['playtime'] for x in f_rows if x['playtime']), ''),
            'popularity': 0,
            'description': '',
            'image': '',
            'image_source': '',
            'youtube': next((x['youtube'] for x in f_rows if x['youtube']), ''),
            'sheet_names': [x['name'] for x in f_rows],
            'sheet_genre': genre_raw,
            'sheet_rating': next((x['rating'] for x in f_rows if x['rating']), ''),
            'sheet_zones': '',
            'sheet_note': '; '.join(dict.fromkeys(x['note'] for x in f_rows if x['note'])),
            'note': '',
        })
        log.append('+ %s: новая строка таблицы, добавлена (%s)' % (name, status))

    games.sort(key=lambda x: x['title'].lower())

    print('Строк в таблице: %d, игр в списке: %d' % (len(rows), len(games)))
    for line in log:
        print('  ' + line)
    if not log:
        print('  изменений нет')
    counts = {}
    for g in games:
        counts[g['status']] = counts.get(g['status'], 0) + 1
    print('Статусы: ' + ', '.join('%s %d' % kv for kv in sorted(counts.items())))
    drafts = [g['title'] for g in games if g['status'] == 'draft']
    if drafts:
        print('Черновики — нужно описание: ' + ', '.join(drafts))

    if args.dry_run:
        print('--dry-run: файлы не изменены')
        return
    io.open(GAMES_JSON, 'w', encoding='utf-8', newline='\n').write(
        json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    subprocess.check_call([sys.executable, os.path.join(ROOT, 'scripts', 'generate-games.py')])


if __name__ == '__main__':
    main()
